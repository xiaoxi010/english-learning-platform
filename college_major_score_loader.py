# college_major_score_loader.py — 各省院校专业录取分数参考（Excel → JSON）
import json
import os
import re

_BASE = os.path.dirname(os.path.abspath(__file__))
_DATA_DIR = os.path.join(_BASE, 'data', 'college_major_scores')

# 表头行关键字 → 字段名（与用户提供的 9 列一致）
_HEADER_ALIASES = {
    '生源地': 'source_region',
    '必选科目': 'required_subjects',
    '科类': 'required_subjects',
    '院校名称': 'school_name',
    '专业名称': 'major_name',
    '专业备注': 'major_note',
    '最低分': 'min_score',
    '最低位次': 'min_rank',
    '最高分': 'max_score',
    '最高位次': 'max_rank',
}

# 导出 JSON 中保留的字段（顺序与展示列一致）
RECORD_FIELDS = [
    'source_region',
    'required_subjects',
    'school_name',
    'major_name',
    'major_note',
    'min_score',
    'min_rank',
    'max_score',
    'max_rank',
]

# 页面展示列（不含生源地；必选科目显示为「选科」）
DISPLAY_COLUMNS = [
    {'key': 'required_subjects', 'label': '选科'},
    {'key': 'school_name', 'label': '大学院校名称'},
    {'key': 'major_name', 'label': '院校专业名称'},
    {'key': 'major_note', 'label': '专业备注'},
    {'key': 'min_score', 'label': '最低分'},
    {'key': 'min_rank', 'label': '最低位'},
    {'key': 'max_score', 'label': '最高分'},
    {'key': 'max_rank', 'label': '最高位'},
]

_DATASETS = {
    'jilin_2026': {
        'province': '吉林',
        'year_label': '2026',
        'title': '吉林省2026院校专业分数参考',
        'xlsx': os.path.join(_DATA_DIR, 'jilin_2026.xlsx'),
        'json': os.path.join(_DATA_DIR, 'jilin_2026.json'),
    },
}


def _norm_header(cell):
    s = re.sub(r'\s+', '', str(cell or '').strip())
    return s


def _cell_str(val):
    if val is None:
        return ''
    if isinstance(val, float) and val == int(val):
        return str(int(val))
    return str(val).strip()


def _cell_num(val):
    if val is None or str(val).strip() == '':
        return None
    try:
        n = float(val)
        return int(n) if n == int(n) else n
    except (TypeError, ValueError):
        return None


def _find_header_row(ws):
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        if not row:
            continue
        labels = [_norm_header(c) for c in row]
        if '院校名称' in labels and '专业名称' in labels:
            return idx, labels
    return None, []


def _build_col_map(labels):
    col_map = {}
    for i, label in enumerate(labels):
        if not label:
            continue
        key = _HEADER_ALIASES.get(label)
        if key and key not in col_map:
            col_map[key] = i
    # 吉林等表：「批次」后一列无表头，值为 历史/物理（科类）→ 必选科目
    if 'required_subjects' not in col_map:
        for i, label in enumerate(labels):
            if label != '批次':
                continue
            j = i + 1
            if j < len(labels) and not labels[j]:
                col_map['required_subjects'] = j
                break
    return col_map


def build_records_from_xlsx(xlsx_path):
    """从 Excel 解析记录，每条包含 9 个展示字段。"""
    if not os.path.isfile(xlsx_path):
        return []

    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    header_idx, labels = _find_header_row(ws)
    if header_idx is None:
        wb.close()
        return []

    col_map = _build_col_map(labels)
    missing = [f for f in RECORD_FIELDS if f not in col_map]
    if missing:
        wb.close()
        raise ValueError(f'Excel 缺少列：{missing}（表头：{labels}）')

    records = []
    for row in ws.iter_rows(min_row=header_idx + 2, values_only=True):
        if not row:
            continue
        school = _cell_str(row[col_map['school_name']]) if col_map['school_name'] < len(row) else ''
        major = _cell_str(row[col_map['major_name']]) if col_map['major_name'] < len(row) else ''
        if not school and not major:
            continue
        rec = {}
        for field in RECORD_FIELDS:
            idx = col_map[field]
            val = row[idx] if idx < len(row) else None
            if field in ('min_score', 'min_rank', 'max_score', 'max_rank'):
                rec[field] = _cell_num(val)
            else:
                rec[field] = _cell_str(val)
        records.append(rec)
    wb.close()
    return records


def build_dataset_json(dataset_id='jilin_2026', xlsx_path=None):
    meta = _DATASETS.get(dataset_id)
    if not meta:
        raise ValueError(f'未知数据集：{dataset_id}')
    path = xlsx_path or meta['xlsx']
    records = build_records_from_xlsx(path)
    payload = {
        'id': dataset_id,
        'province': meta['province'],
        'year_label': meta['year_label'],
        'title': meta['title'],
        'columns': DISPLAY_COLUMNS,
        'count': len(records),
        'records': records,
    }
    out = meta['json']
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, 'w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False, indent=None, separators=(',', ':'))
    return payload


def load_dataset(dataset_id='jilin_2026'):
    meta = _DATASETS.get(dataset_id)
    if not meta:
        return None
    json_path = meta['json']
    if os.path.isfile(json_path):
        with open(json_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        data['columns'] = DISPLAY_COLUMNS
        return data
    if os.path.isfile(meta['xlsx']):
        return build_dataset_json(dataset_id)
    return None


def search_records(dataset_id='jilin_2026', query='', school='', major='', limit=80):
    data = load_dataset(dataset_id)
    if not data:
        return {'ok': False, 'message': '数据文件不存在', 'records': [], 'total': 0}

    school_q = (school or '').strip().lower()
    major_q = (major or '').strip().lower()
    legacy_q = (query or '').strip().lower()

    if legacy_q and not school_q and not major_q:
        school_q = legacy_q
        major_q = legacy_q

    records = data.get('records') or []
    if not school_q and not major_q:
        return {
            'ok': True,
            'total': 0,
            'records': [],
            'columns': data.get('columns'),
            'title': data.get('title'),
        }

    def _bits(text):
        return [b for b in re.split(r'[\s,，、]+', text) if b]

    school_bits = _bits(school_q) if school_q else []
    major_bits = _bits(major_q) if major_q else []
    matched = []
    for rec in records:
        if school_bits:
            name = str(rec.get('school_name') or '').lower()
            if not all(bit in name for bit in school_bits):
                continue
        if major_bits:
            blob = ' '.join([
                str(rec.get('major_name') or ''),
                str(rec.get('major_note') or ''),
            ]).lower()
            if not all(bit in blob for bit in major_bits):
                continue
        matched.append(rec)
        if len(matched) >= limit:
            break

    return {
        'ok': True,
        'total': len(matched),
        'records': matched,
        'columns': data.get('columns'),
        'title': data.get('title'),
    }


def _no_data_record():
    rec = {}
    for field in RECORD_FIELDS:
        rec[field] = '暂无数据'
    return rec


def lookup_record(dataset_id='jilin_2026', school='', major='', subject_track=''):
    """按院校 + 专业查一条分数参考（优先匹配科类）。"""
    data = load_dataset(dataset_id)
    if not data:
        return {'ok': False, 'message': '数据文件不存在', 'found': False, 'record': _no_data_record()}
    school_q = (school or '').strip()
    major_q = (major or '').strip()
    if not school_q or not major_q:
        return {'ok': True, 'found': False, 'record': _no_data_record()}

    result = search_records(dataset_id, school=school_q, major=major_q, limit=50)
    matched = result.get('records') or []
    exact = [r for r in matched if str(r.get('major_name') or '').strip() == major_q]
    if exact:
        matched = exact
    track = (subject_track or '').strip()
    if track and matched:
        filtered = [r for r in matched if str(r.get('required_subjects') or '') == track]
        if filtered:
            matched = filtered
    if not matched:
        return {'ok': True, 'found': False, 'record': _no_data_record()}
    return {'ok': True, 'found': True, 'record': matched[0]}


def list_schools(dataset_id='jilin_2026', query='', limit=40):
    data = load_dataset(dataset_id)
    if not data:
        return {'ok': False, 'schools': []}
    q = (query or '').strip().lower()
    seen = set()
    schools = []
    for rec in data.get('records') or []:
        name = str(rec.get('school_name') or '').strip()
        if not name or name in seen:
            continue
        if q and q not in name.lower():
            continue
        seen.add(name)
        schools.append(name)
        if len(schools) >= limit:
            break
    schools.sort()
    return {'ok': True, 'schools': schools}


def dataset_status(dataset_id='jilin_2026'):
    meta = _DATASETS.get(dataset_id)
    if not meta:
        return {'available': False}
    has_xlsx = os.path.isfile(meta['xlsx'])
    has_json = os.path.isfile(meta['json'])
    count = 0
    if has_json:
        try:
            with open(meta['json'], 'r', encoding='utf-8') as f:
                count = json.load(f).get('count') or 0
        except (json.JSONDecodeError, OSError):
            count = 0
    return {
        'available': has_xlsx or has_json,
        'has_xlsx': has_xlsx,
        'has_json': has_json,
        'count': count,
        **{k: meta[k] for k in ('province', 'year_label', 'title')},
    }
